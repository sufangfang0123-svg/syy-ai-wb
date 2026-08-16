import io

from docx import Document
from sqlalchemy import select


def enable_fixture(monkeypatch, mode="success"):
    monkeypatch.setenv("NDG_ENVIRONMENT", "test")
    monkeypatch.setenv("NDG_AI_FIXTURE_PROVIDER", "1")
    monkeypatch.setenv("NDG_AI_FIXTURE_MODE", mode)


def create_source(client, project, text="样本记录显示，12名受试者中有9名选择可重复收纳包装。"):
    response = client.post(f"/api/v1/projects/{project['id']}/ai/sources/text", json={"source_name": "固定脱敏评测材料", "text": text})
    assert response.status_code == 201
    return response.json()


def create_run(client, project, source):
    return client.post(f"/api/v1/projects/{project['id']}/ai/evidence-runs", json={"source_document_id": source["id"]})


def test_provider_disabled_is_explicit_and_manual_flow_survives(client, project, monkeypatch):
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    monkeypatch.delenv("OPENAI_MODEL", raising=False)
    monkeypatch.delenv("NDG_AI_FIXTURE_PROVIDER", raising=False)
    status = client.get("/api/v1/ai/evidence-copilot/status").json()
    assert status["configured"] is False and status["provider"] == "disabled"
    source = create_source(client, project)
    failed = create_run(client, project, source)
    assert failed.status_code == 503 and "未配置" in failed.json()["detail"]
    manual = client.post(f"/api/v1/projects/{project['id']}/evidence/manual", json={"title": "手工证据", "raw_text": "手工流程继续可用"})
    assert manual.status_code == 201


def test_ai_run_creates_only_review_candidates_without_revision_or_gate_change(client, project, monkeypatch):
    enable_fixture(monkeypatch)
    gate = client.post(f"/api/v1/projects/{project['id']}/gate").json()
    source = create_source(client, project)
    assert client.get(f"/api/v1/projects/{project['id']}").json()["revision"] == 1
    response = create_run(client, project, source)
    assert response.status_code == 201
    run = response.json()
    assert run["status"] == "SUCCEEDED" and len(run["candidates"]) == 1
    candidate = run["candidates"][0]
    assert candidate["citation_verification_status"] == "VERIFIED"
    assert candidate["review_status"] == "PENDING_REVIEW"
    assert client.get(f"/api/v1/projects/{project['id']}").json()["revision"] == 1
    assert client.get(f"/api/v1/projects/{project['id']}/gate/current").json()["id"] == gate["id"]
    assert client.get(f"/api/v1/projects/{project['id']}/evidence").json() == []


def test_accept_edit_reject_idempotency_revision_stale_audit_and_export(client, project, monkeypatch):
    enable_fixture(monkeypatch)
    gate = client.post(f"/api/v1/projects/{project['id']}/gate").json()
    first = create_run(client, project, create_source(client, project)).json()["candidates"][0]
    headers = {"Idempotency-Key": "accept-once-0001", "If-Match": "1"}
    accepted = client.post(f"/api/v1/projects/{project['id']}/ai/evidence-candidates/{first['id']}/review", headers=headers, json={"action": "ACCEPT", "reviewer_note": "负责人逐字核对"})
    assert accepted.status_code == 200
    assert accepted.json()["review_status"] == "ACCEPTED" and accepted.json()["final_evidence_id"]
    assert client.get(f"/api/v1/projects/{project['id']}").json()["revision"] == 2
    assert client.get(f"/api/v1/projects/{project['id']}/gate/current").status_code == 404
    repeated = client.post(f"/api/v1/projects/{project['id']}/ai/evidence-candidates/{first['id']}/review", headers={"Idempotency-Key": "accept-once-0001", "If-Match": "2"}, json={"action": "ACCEPT"})
    assert repeated.status_code == 200
    assert client.get(f"/api/v1/projects/{project['id']}").json()["revision"] == 2

    second_source = create_source(client, project, "另一段材料明确记录，样品在两次受控测试中均保持完整。")
    second = create_run(client, project, second_source).json()["candidates"][0]
    edited = client.post(f"/api/v1/projects/{project['id']}/ai/evidence-candidates/{second['id']}/review", headers={"Idempotency-Key": "edited-once-0002", "If-Match": "2"}, json={"action": "EDIT_AND_ACCEPT", "claim": "受控材料记录样品在两次测试中保持完整。", "scope": "仅适用于这两次受控测试。", "limitations": "不能外推长期质量。", "reviewer_note": "负责人收窄主张"})
    assert edited.status_code == 200 and edited.json()["review_status"] == "EDITED_AND_ACCEPTED"

    third_source = create_source(client, project, "第三段材料仅用于验证拒绝流程，不形成正式证据。")
    third = create_run(client, project, third_source).json()["candidates"][0]
    rejected = client.post(f"/api/v1/projects/{project['id']}/ai/evidence-candidates/{third['id']}/review", headers={"Idempotency-Key": "reject-once-0003", "If-Match": "3"}, json={"action": "REJECT", "reviewer_note": "范围不足"})
    assert rejected.status_code == 200 and rejected.json()["review_status"] == "REJECTED"
    assert client.get(f"/api/v1/projects/{project['id']}").json()["revision"] == 3

    exported = client.get(f"/api/v1/projects/{project['id']}/export").json()
    assert len(exported["evidence"]) == 2
    provenance = exported["ai_provenance"]
    assert len(provenance["runs"]) == 3
    assert {item["review_status"] for run in provenance["runs"] for item in run["candidates"]} == {"ACCEPTED", "EDITED_AND_ACCEPTED", "REJECTED"}
    actions = {item["action"] for item in exported["audit_events"]}
    assert {"AI_SOURCE_CREATED", "AI_RUN_STARTED", "AI_RUN_SUCCEEDED", "AI_CANDIDATE_ACCEPTED", "AI_CANDIDATE_EDITED_ACCEPTED", "AI_CANDIDATE_REJECTED"}.issubset(actions)
    assert gate["rule_version"] == "NDG_GATE_V0.3.0"


def test_provider_failures_are_persisted_without_fake_success(client, project, monkeypatch):
    cases = (
        ("timeout", 504, "TIMED_OUT", "provider_timeout"),
        ("rate_limit", 503, "FAILED", "provider_rate_limited"),
        ("network", 503, "FAILED", "provider_network_error"),
        ("refusal", 422, "ABSTAINED", "provider_refusal"),
        ("malformed", 422, "FAILED", "malformed_output"),
    )
    for index, (mode, status_code, stored_status, error_code) in enumerate(cases, 1):
        enable_fixture(monkeypatch, mode)
        source = create_source(client, project, f"失败模式{index}的固定脱敏材料。")
        response = create_run(client, project, source)
        assert response.status_code == status_code
        runs = client.get(f"/api/v1/projects/{project['id']}/ai/evidence-runs").json()
        assert runs[0]["status"] == stored_status and runs[0]["sanitized_error_code"] == error_code
        assert runs[0]["candidates"] == []
    assert client.get(f"/api/v1/projects/{project['id']}/evidence").json() == []
    assert client.get(f"/api/v1/projects/{project['id']}").json()["revision"] == 1


def test_provider_http_errors_are_safely_classified_without_raw_details():
    from app.ai import map_provider_status_error

    expected = {
        401: "provider_auth_error",
        403: "provider_auth_error",
        404: "provider_model_unavailable",
        400: "provider_request_rejected",
        422: "provider_request_rejected",
        500: "provider_error",
    }
    for status, code in expected.items():
        failure = map_provider_status_error(status)
        assert failure.code == code and failure.http_status == 503
        assert "sk-" not in failure.public_message and "Traceback" not in failure.public_message


def test_invalid_and_manual_citations_are_blocked(client, project, monkeypatch):
    enable_fixture(monkeypatch)
    source = create_source(client, project)
    candidate = create_run(client, project, source).json()["candidates"][0]
    # API context never treats text presence as source truth.
    context = client.get(f"/api/v1/projects/{project['id']}/ai/evidence-candidates/{candidate['id']}/context").json()
    assert "不证明材料本身正确" in context["warning"]
    from app.copilot_service import verify_citation
    from app.materials import segment_plain_text
    text, locators = segment_plain_text("第一段事实。\n\n第二段相反事实。")
    assert verify_citation(text, locators, "第一段事实。", "paragraph-0002")[0] == "NEEDS_MANUAL_VERIFICATION"
    assert verify_citation(text, locators, "并不存在的引用。", "paragraph-0001")[0] == "INVALID"


def test_docx_source_and_duplicate_sha(client, project):
    output = io.BytesIO()
    document = Document()
    document.add_paragraph("第一段可定位的脱敏材料。")
    document.add_paragraph("第二段包含有限事实。")
    document.save(output)
    payload = output.getvalue()
    response = client.post(f"/api/v1/projects/{project['id']}/ai/sources/file", files={"file": ("fixture.docx", payload, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 201 and response.json()["extraction_status"] == "READY"
    duplicate = client.post(f"/api/v1/projects/{project['id']}/ai/sources/file", files={"file": ("fixture.docx", payload, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert duplicate.status_code == 409


def test_interrupted_run_recovery(tmp_path):
    from app.copilot_service import mark_interrupted_runs
    from app.database import make_engine, run_migrations, sessionmaker
    from app.models import AIEvidenceRun, AISourceDocument, Project
    engine = make_engine(f"sqlite:///{(tmp_path / 'restart.sqlite3').as_posix()}")
    run_migrations(engine)
    Session = sessionmaker(bind=engine, expire_on_commit=False)
    with Session.begin() as session:
        project = Project(name="重启恢复", decision_question="能否恢复未完成AI运行？")
        session.add(project)
        session.flush()
        source = AISourceDocument(project_id=project.id, source_kind="text", source_name="夹具", mime_type="text/plain", sha256="a" * 64, extracted_text="正文", locator_map="[]")
        session.add(source)
        session.flush()
        session.add(AIEvidenceRun(project_id=project.id, source_document_id=source.id, provider="fixture", model="fixture", prompt_version="P", output_schema_version="S", status="RUNNING"))
    engine.dispose()
    restarted = make_engine(f"sqlite:///{(tmp_path / 'restart.sqlite3').as_posix()}")
    Restarted = sessionmaker(bind=restarted, expire_on_commit=False)
    with Restarted() as session:
        assert mark_interrupted_runs(session) == 1
        assert session.scalar(select(AIEvidenceRun)).status == "INTERRUPTED"
