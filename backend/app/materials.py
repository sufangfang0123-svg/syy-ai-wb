from __future__ import annotations

import csv
import hashlib
import io
import os
import zipfile
from pathlib import Path, PurePath

from docx import Document
from pypdf import PdfReader

MAX_FILE_BYTES = 5 * 1024 * 1024
ALLOWED_TYPES = {
    ".txt": {"text/plain"},
    ".md": {"text/markdown", "text/plain"},
    ".csv": {"text/csv", "application/csv", "text/plain"},
    ".pdf": {"application/pdf"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"},
}
MAX_EXTRACTED_CHARS = 1_000_000
MAX_PDF_PAGES = 50
MAX_DOCX_UNCOMPRESSED_BYTES = 20 * 1024 * 1024


def extract_upload(filename: str, declared_mime: str | None, data: bytes) -> dict[str, str | int]:
    result = extract_material(filename, declared_mime, data, allow_manual_verification=False)
    return {key: value for key, value in result.items() if key not in {"locator_map", "extraction_status"}}  # type: ignore[return-value]


def extract_material(filename: str, declared_mime: str | None, data: bytes, *, allow_manual_verification: bool = True) -> dict[str, object]:
    if not filename or PurePath(filename).name != filename or "/" in filename or "\\" in filename:
        raise ValueError("文件名包含非法路径")
    if any(character in filename for character in '<>:"|?*\x00') or filename.endswith((".", " ")):
        raise ValueError("文件名包含非法字符")
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_TYPES:
        raise ValueError("仅支持 txt、md、csv、docx 和文本型 PDF")
    mime = (declared_mime or "").split(";", 1)[0].lower()
    if mime not in ALLOWED_TYPES[suffix]:
        raise ValueError("文件扩展名与 MIME 类型不匹配")
    if not data or len(data) > MAX_FILE_BYTES:
        raise ValueError("文件必须大于 0 且不超过 5MB")
    digest = hashlib.sha256(data).hexdigest()
    locator_map: list[dict[str, object]] = []
    extraction_status = "READY"
    if suffix == ".pdf":
        if not data.startswith(b"%PDF-"):
            raise ValueError("PDF 文件签名无效")
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("不接受加密 PDF")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError(f"PDF页数超过{MAX_PDF_PAGES}页限制")
            segments = []
            for page_number, page in enumerate(reader.pages, 1):
                page_text = (page.extract_text() or "").strip()
                if page_text:
                    segments.append((f"page-{page_number:04d}", f"第{page_number}页", page_text, {"page": page_number}))
            text, locator_map = _assemble_segments(segments)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("PDF 解析失败") from exc
        if not text:
            if not allow_manual_verification:
                raise ValueError("PDF 无可提取文本；当前不支持扫描件 OCR")
            extraction_status = "NEEDS_MANUAL_VERIFICATION"
    elif suffix == ".docx":
        if not data.startswith(b"PK"):
            raise ValueError("DOCX 文件签名无效")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = archive.namelist()
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise ValueError("DOCX 文件结构无效")
                if len(names) > 1_000 or sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise ValueError("DOCX 解压后内容超过安全限制")
            document = Document(io.BytesIO(data))
            segments = []
            ordinal = 0
            for paragraph in document.paragraphs:
                value = paragraph.text.strip()
                if value:
                    ordinal += 1
                    segments.append((f"paragraph-{ordinal:04d}", f"第{ordinal}段", value, {"paragraph": ordinal}))
            for table_index, table in enumerate(document.tables, 1):
                for row_index, row in enumerate(table.rows, 1):
                    value = " | ".join(cell.text.strip() for cell in row.cells).strip(" |")
                    if value:
                        ordinal += 1
                        segments.append((f"table-{table_index:03d}-row-{row_index:04d}", f"表{table_index}第{row_index}行", value, {"paragraph": ordinal}))
            text, locator_map = _assemble_segments(segments)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("DOCX 解析失败") from exc
        if not text:
            if not allow_manual_verification:
                raise ValueError("DOCX 无可提取文本")
            extraction_status = "NEEDS_MANUAL_VERIFICATION"
    else:
        if b"\x00" in data:
            raise ValueError("文本文件包含二进制空字符")
        try:
            decoded = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise ValueError("文本文件必须使用 UTF-8 编码") from exc
        if suffix == ".csv":
            rows = list(csv.reader(io.StringIO(decoded)))
            if not rows:
                raise ValueError("CSV 为空")
            segments = [(f"row-{index:05d}", f"CSV第{index}行", " | ".join(cell.strip() for cell in row), {"paragraph": index}) for index, row in enumerate(rows, 1)]
        else:
            paragraphs = [value.strip() for value in decoded.replace("\r\n", "\n").replace("\r", "\n").split("\n\n") if value.strip()]
            segments = [(f"paragraph-{index:04d}", f"第{index}段", value, {"paragraph": index}) for index, value in enumerate(paragraphs, 1)]
        text, locator_map = _assemble_segments(segments)
        if not text:
            raise ValueError("文件没有可导入文本")
    if len(text) > MAX_EXTRACTED_CHARS:
        raise ValueError("提取文本超过 100 万字符")
    return {"raw_text": text, "sha256": digest, "mime_type": mime, "size_bytes": len(data), "suffix": suffix, "locator_map": locator_map, "extraction_status": extraction_status}


def segment_plain_text(text: str) -> tuple[str, list[dict[str, object]]]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    if not normalized:
        raise ValueError("材料正文不能为空")
    if len(normalized) > MAX_EXTRACTED_CHARS:
        raise ValueError("材料正文超过 100 万字符")
    paragraphs = [value.strip() for value in normalized.split("\n\n") if value.strip()]
    segments = [(f"paragraph-{index:04d}", f"第{index}段", value, {"paragraph": index}) for index, value in enumerate(paragraphs, 1)]
    return _assemble_segments(segments)


def _assemble_segments(segments: list[tuple[str, str, str, dict[str, int]]]) -> tuple[str, list[dict[str, object]]]:
    parts: list[str] = []
    locator_map: list[dict[str, object]] = []
    cursor = 0
    for segment_id, label, value, metadata in segments:
        cleaned = value.strip()
        if not cleaned:
            continue
        if parts:
            parts.append("\n\n")
            cursor += 2
        start = cursor
        parts.append(cleaned)
        cursor += len(cleaned)
        locator_map.append({"segment_id": segment_id, "label": label, "start": start, "end": cursor, **metadata})
    return "".join(parts), locator_map


def store_snapshot(project_id: str, digest: str, suffix: str, data: bytes) -> str:
    data_root = Path(os.getenv("NDG_DATA_DIR", Path(__file__).resolve().parents[2] / "data")).resolve()
    # Keep the physical path short enough for Windows hosts with legacy MAX_PATH
    # handling. The full project id and SHA-256 remain in SQLite; snapshot_ref is
    # only an opaque local storage locator.
    project_bucket = hashlib.sha256(project_id.encode("utf-8")).hexdigest()[:16]
    snapshot_name = f"{digest[:32]}{suffix}.snapshot"
    upload_root = (data_root / "uploads" / project_bucket).resolve()
    if data_root not in upload_root.parents:
        raise ValueError("快照路径越界")
    upload_root.mkdir(parents=True, exist_ok=True)
    target = (upload_root / snapshot_name).resolve()
    if upload_root not in target.parents:
        raise ValueError("快照路径越界")
    if not target.exists():
        with target.open("xb") as handle:
            handle.write(data)
    return target.relative_to(data_root).as_posix()
